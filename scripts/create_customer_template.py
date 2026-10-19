#!/usr/bin/env python3
"""Create an empty Google Sheets import template matching the Apps Script schema."""
import ast
import json
from pathlib import Path
import re

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'templates/ok-express-customer-list.xlsx'
NAVY, ORANGE, PAPER, INK = '15283E', 'CA4715', 'F7F6F2', '172B40'


def style_table(sheet, headers, widths, freeze='A2'):
    sheet.append(headers)
    sheet.freeze_panes = freeze
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.sheet_properties.tabColor = NAVY
    sheet.row_dimensions[1].height = 42
    sheet.auto_filter.ref = f'A1:{get_column_letter(len(headers))}1001'
    for column, (label, width) in enumerate(zip(headers, widths), 1):
        letter = get_column_letter(column)
        sheet.column_dimensions[letter].width = width
        cell = sheet.cell(1, column)
        cell.fill = PatternFill('solid', fgColor=NAVY)
        cell.font = Font(name='맑은 고딕', size=11, color='FFFFFF', bold=True)
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        for row in range(2, 1002):
            cell = sheet.cell(row, column)
            cell.font = Font(name='맑은 고딕', size=10, color=INK)
            cell.fill = PatternFill('solid', fgColor='FFFFFF' if row % 2 == 0 else 'F5F7F8')
            cell.number_format = '@'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    sheet.print_options.horizontalCentered = True
    sheet.sheet_properties.outlinePr.summaryRight = False


def dropdown(sheet, area, options):
    validation = DataValidation(type='list', formula1='"' + ','.join(options) + '"', allow_blank=True)
    validation.errorTitle = '목록에서 선택해 주세요'
    validation.error = '입력값을 확인해 주세요.'
    validation.showErrorMessage = True
    sheet.add_data_validation(validation)
    validation.add(area)


def build():
    config = json.loads((ROOT / 'config.json').read_text(encoding='utf-8'))
    source = (ROOT / 'google-apps-script/Code.gs').read_text(encoding='utf-8')
    headers = ast.literal_eval(re.search(r'var HEADERS = (\[.*?\]);', source).group(1))
    workbook = Workbook()
    workbook.properties.creator = config['business']['name']
    workbook.properties.title = 'OK익스프레스 고객 문의 명단'
    workbook.properties.description = '빈 고객명단 템플릿. 실제 Google 연결·메일 발송은 별도 설정이 필요합니다.'
    sheet = workbook.active
    sheet.title = '이사 문의'
    style_table(sheet, headers, [26, 23, 21, 15, 25, 25, 19, 16, 13, 45, 16, 16, 17, 28, 36, 16, 23, 25], 'D2')
    sheet.column_dimensions['R'].hidden = True
    sheet['A1'].comment = Comment('Netlify 접수 ID입니다. 중복 처리에 사용하므로 변경하지 마세요.', 'OK익스프레스')
    sheet['C1'].comment = Comment('텍스트 형식으로 전화번호 앞자리 0을 유지합니다.', 'OK익스프레스')
    sheet['P1'].comment = Comment('PENDING: 전송 대기, SENT: Google 메일 전송 호출 성공, FAILED: 전송 실패. SENT도 실제 수신함 도착을 보장하지 않습니다. 상태를 임의로 변경하지 마세요.', 'OK익스프레스')
    sheet['R1'].comment = Comment('중복·변경 확인용 시스템 값입니다. 숨김 상태를 유지하고 편집하지 마세요.', 'OK익스프레스')
    dropdown(sheet, 'D2:D1001', ['포장이사', '일반이사', '사무실이사'])
    for status, color, text in [('SENT', 'E6F1EB', '236241'), ('FAILED', 'FCE8E1', 'A73612'), ('PENDING', 'FFF3D1', '785C0C')]:
        sheet.conditional_formatting.add('P2:P1001', CellIsRule(operator='equal', formula=['"' + status + '"'], fill=PatternFill('solid', fgColor=color), font=Font(color=text)))

    management = workbook.create_sheet('상담 관리')
    style_table(management, ['접수 ID', '상담 상태', '담당자', '최종 연락일', '다음 연락일', '상담 메모', '목적 달성일', '파기 확인일'], [28, 17, 17, 19, 19, 50, 19, 19], 'C2')
    management.sheet_properties.tabColor = ORANGE
    dropdown(management, 'B2:B1001', ['미연락', '연락시도', '상담중', '견적안내', '계약전환', '상담종료'])
    for column in ['D', 'E', 'G', 'H']:
        for row in range(2, 1002):
            management[f'{column}{row}'].number_format = 'yyyy-mm-dd'
    management['A1'].comment = Comment('이사 문의 탭의 접수 ID를 복사해 상담 현황을 수동 관리하세요. 이 탭은 자동으로 채워지지 않습니다.', 'OK익스프레스')
    management['F1'].comment = Comment('상담에 필요한 내용만 남기세요. 주민등록번호·계좌번호·상세주소·출입 비밀번호 등은 적지 마세요.', 'OK익스프레스')

    guide = workbook.create_sheet('사용 안내')
    guide.sheet_view.showGridLines = False
    guide.sheet_properties.tabColor = '7B8793'
    for col in range(1, 9): guide.column_dimensions[get_column_letter(col)].width = 15
    guide.merge_cells('A1:H2')
    guide['A1'] = 'OK익스프레스 | 고객 문의 명단'
    guide['A1'].font = Font(name='맑은 고딕', size=22, bold=True, color='FFFFFF')
    guide['A1'].fill = PatternFill('solid', fgColor=NAVY)
    guide['A1'].alignment = Alignment(vertical='center')
    guide.row_dimensions[1].height = 26
    guide.row_dimensions[2].height = 24
    notes = [
        ('현재 상태', '빈 명단 템플릿입니다. 실제 고객 정보·예시 고객은 넣지 않았습니다. Google 계정 안에 시트를 생성하거나 실제 연동을 완료한 상태는 아닙니다.'),
        ('구글 시트로 가져오기', 'Google 스프레드시트에서 파일 → 가져오기 → 업로드 → 이 XLSX 선택 → 새 스프레드시트 만들기를 선택하세요. 가져온 뒤 시트 이름과 첫 행을 확인하세요.'),
        ('이사 문의 탭', '사이트 연동으로 들어오는 원본입니다. 탭 이름과 A~R열의 제목·순서를 바꾸지 마세요. 첫 행은 Apps Script 코드의 HEADERS와 일치합니다.'),
        ('상담 관리 탭', '선택적으로 사용하는 수동 관리 탭입니다. 접수 ID를 복사한 뒤 상태·담당자·연락일을 직접 기록합니다. 자동 동기화나 자동 알림 기능은 없습니다.'),
        ('Apps Script 연결', '가져온 Google 시트에서 확장 프로그램 → Apps Script를 열고 저장소의 Code.gs와 manifest를 적용한 뒤 setup을 실행하세요. setup이 본인 Google 이메일을 수신자로 설정합니다.'),
        ('계정 설정 안내', 'https://github.com/rhwnssla133-jpg/-/blob/main/docs/GOOGLE-SHEETS.md'),
        ('연동 코드', 'https://github.com/rhwnssla133-jpg/-/blob/main/google-apps-script/Code.gs'),
        ('배포 방식', 'Google 자동 연동은 Netlify의 GitHub 연결 배포와 Functions 설정이 필요합니다. 정적 ZIP만 업로드해서는 서버 연동 함수가 배포되지 않습니다.'),
        ('알림 상태', 'PENDING: 대기 / SENT: 메일 전송 호출 성공 / FAILED: 실패. 실제 수신함과 스팸함도 확인하세요. 실패 재전송은 시트의 OK익스프레스 연동 메뉴에서 실행합니다.'),
        ('선택정보', '선택 수집 동의가 없으면 호칭·날짜·요청사항은 전송하지 않습니다. 필요한 상담 정보만 기록하고 민감하거나 불필요한 정보를 넣지 마세요.'),
        ('공유와 비밀키', '시트는 비공개로 두고 필요한 담당자만 공유합니다. 비밀키·비밀번호·인증 코드를 이 파일이나 GitHub에 저장하지 마세요.'),
        ('보유기간', '상담 목적 달성 시 지체 없이 파기하고, 미계약 문의는 최대 90일 보관하는 운영안입니다. 법정기간이나 자동 삭제 기능을 뜻하지 않습니다.'),
        ('삭제 범위', 'Netlify 원본, 이사 문의 행, 상담 관리 기록, 이메일, 휴지통과 별도 사본을 함께 관리하세요. Google 버전 기록·백업의 실제 보유 조건도 확인해야 합니다.'),
        ('접수 개시 전', '사업정보·개인정보 담당자·Google/Netlify 국외이전 안내와 동의, 실제 접수·행 생성·메일 수신을 확인한 뒤 온라인 접수를 여세요.'),
    ]
    for row, (title, note) in enumerate(notes, 4):
        guide.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
        guide.cell(row, 1, title)
        guide.cell(row, 2, note)
        guide.row_dimensions[row].height = 56 if len(note) > 110 else 44
        for col in [1, 2]:
            cell = guide.cell(row, col)
            cell.font = Font(name='맑은 고딕', size=10, color=INK, bold=col == 1)
            cell.fill = PatternFill('solid', fgColor=PAPER if row % 2 == 0 else 'FFFFFF')
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        if note.startswith('https://'):
            guide.cell(row, 2).hyperlink = note
            guide.cell(row, 2).font = Font(name='맑은 고딕', size=10, color='1260C4', underline='single')
    guide.freeze_panes = 'B4'
    OUTPUT.parent.mkdir(exist_ok=True)
    workbook.save(OUTPUT)

    # Confirm the actual saved workbook preserves the live integration contract.
    saved = load_workbook(OUTPUT)
    assert saved.sheetnames == ['이사 문의', '상담 관리', '사용 안내']
    assert [cell.value for cell in saved['이사 문의'][1]] == headers
    for tab in ['이사 문의', '상담 관리']:
        assert all(cell.value is None for row in saved[tab].iter_rows(min_row=2) for cell in row)
    assert saved['이사 문의']['C2'].number_format == '@'
    assert saved['이사 문의'].column_dimensions['R'].hidden
    print(f'Created {OUTPUT}; exact {len(headers)} integration columns; 0 customer records; 3 sheets')


if __name__ == '__main__':
    build()
